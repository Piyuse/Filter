"""Bounded, local parsers. Uploaded originals are never saved by this app."""
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, BadZipFile
import os
import shutil
import warnings
from PIL import Image
import pytesseract
from pypdf import PdfReader
from docx import Document
from .errors import UserInputError
from .ocr import recognize, diagnostics
from .pdf_text import extract_pdf_page
from .text_format import format_document_text

MAX_BYTES = 10 * 1024 * 1024
MAX_PAGES = 20
Image.MAX_IMAGE_PIXELS = 20_000_000

def find_tesseract():
    """Use a configured executable or a local install without relying on stale PATH."""
    configured = os.environ.get('TESSERACT_CMD')
    if configured:
        return configured
    local = Path(__file__).resolve().parent.parent / '.tools' / 'tesseract' / 'tesseract.exe'
    if local.is_file():
        return str(local)
    executable = shutil.which('tesseract')
    if executable:
        return executable
    if os.name == 'nt':
        roots = [Path(os.environ.get('ProgramFiles', 'C:/Program Files')),
                 Path(os.environ.get('ProgramFiles(x86)', 'C:/Program Files (x86)')),
                 Path(os.environ.get('LOCALAPPDATA', str(Path.home() / 'AppData/Local'))) / 'Programs']
        for root in roots:
            candidate = root / 'Tesseract-OCR' / 'tesseract.exe'
            if candidate.is_file():
                return str(candidate)
    return 'tesseract'

def ocr_available():
    # Resolve on each check so installing OCR does not require a terminal or PC
    # restart and the health route uses the same executable as upload processing.
    pytesseract.pytesseract.tesseract_cmd = find_tesseract()
    return bool(shutil.which(pytesseract.pytesseract.tesseract_cmd))

def image_text(image, details=None):
    if not ocr_available():
        raise UserInputError('OCR needs local Tesseract. Install it and set TESSERACT_CMD; text-based PDF, DOCX and TXT still work.')
    if image.width * image.height > Image.MAX_IMAGE_PIXELS:
        raise UserInputError('Image exceeds the 20 megapixel limit.')
    result = recognize(image)
    if details is not None:
        details.append(result)
    return result.text

def extract_text(file, *, metadata=None):
    ocr_details = []
    suffix = Path(file.filename or '').suffix.lower()
    if suffix not in ('.txt', '.docx', '.pdf', '.png', '.jpg', '.jpeg', '.tif', '.tiff', '.webp'):
        raise UserInputError('Supported files: PDF, DOCX, TXT, PNG, JPG, TIFF and WebP.')
    raw = file.stream.read(MAX_BYTES + 1)
    if len(raw) > MAX_BYTES:
        raise UserInputError('Files must be smaller than 10 MB.')
    if not raw:
        raise UserInputError('This file is empty.')
    stream = BytesIO(raw)
    if suffix == '.txt':
        try:
            text = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            raise UserInputError('TXT files must use UTF-8 encoding.') from None
    elif suffix == '.docx':
        try:
            with ZipFile(stream) as archive:
                if len(archive.infolist()) > 2000 or sum(i.file_size for i in archive.infolist()) > 30 * 1024 * 1024:
                    raise UserInputError('DOCX contents exceed the safe extraction limit.')
                if 'word/document.xml' not in archive.namelist():
                    raise UserInputError('Invalid DOCX document.')
            stream.seek(0)
            doc = Document(stream)
            parts = [p.text for p in doc.paragraphs]
            parts += [' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows]
            for section in doc.sections:
                parts += [p.text for p in section.header.paragraphs + section.footer.paragraphs]
            text = '\n'.join(parts)
        except BadZipFile:
            raise UserInputError('Invalid DOCX document.') from None
    elif suffix == '.pdf':
        reader = PdfReader(stream)
        if reader.is_encrypted:
            raise UserInputError('Password-protected PDFs are not supported.')
        if len(reader.pages) > MAX_PAGES:
            raise UserInputError('PDFs are limited to 20 pages.')
        parts = []
        raster = None
        for index, page in enumerate(reader.pages):
            content = page.get_contents()
            if content and len(content.get_data()) > 20 * 1024 * 1024:
                raise UserInputError('PDF page exceeds the safe extraction limit.')
            part = extract_pdf_page(page)
            if not part.strip():
                import pypdfium2 as pdfium
                if raster is None:
                    raster = pdfium.PdfDocument(raw)
                pdf_page = raster[index]
                if pdf_page.get_width() * pdf_page.get_height() * 4 > Image.MAX_IMAGE_PIXELS:
                    raise UserInputError('Scanned PDF page exceeds the image size limit.')
                bitmap = pdf_page.render(scale=2)
                try:
                    part = image_text(bitmap.to_pil(), ocr_details)
                finally:
                    bitmap.close()
                    pdf_page.close()
            parts.append(part)
        if raster:
            raster.close()
        text = '\n\n'.join(parts)
    else:
        with warnings.catch_warnings():
            warnings.simplefilter('error', Image.DecompressionBombWarning)
            with Image.open(stream) as image:
                if getattr(image, 'n_frames', 1) > MAX_PAGES:
                    raise UserInputError('Images are limited to 20 frames/pages.')
                parts = []
                for frame in range(getattr(image, 'n_frames', 1)):
                    image.seek(frame)
                    parts.append(image_text(image, ocr_details))
                text = '\n\n'.join(parts)
    if suffix != '.txt':
        text = format_document_text(text)
    if not text.strip():
        raise UserInputError('No readable text found. Check the document or OCR language settings.')
    if len(text) > 100000:
        raise UserInputError('Extracted text exceeds the 100,000 character limit.')
    if metadata is not None:
        metadata.update(diagnostics(ocr_details))
    return text
