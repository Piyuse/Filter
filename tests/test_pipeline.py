from io import BytesIO
from socket import socket
import pytest
from docx import Document
from pypdf import PdfReader
from reportlab.pdfgen import canvas
from app.main import create_app
from app.redactor import anonymize
from app.language import detect_language

@pytest.fixture
def client():
    return create_app({'TESTING': True}).test_client()

TEXT = '''Patient: Olivia Bennett
Date of birth: 08/14/1986
SSN: 482-91-7364
MRN: MR-209481
Email: olivia.bennett@example.com
Phone: +1 (415) 555-0198
Address: 2847 Maple Street, San Francisco, CA 94110
Visit date: October 22, 2026
Dr. James Wilson
Lab accession: LAB-2026-08432
Olivia Bennett reported fatigue. Blood pressure was 118/76 mmHg.'''

def test_sensitive_values_removed_and_clinical_context_preserved():
    result = anonymize(TEXT)
    for value in ('Olivia Bennett', '08/14/1986', '482-91-7364', 'MR-209481',
                  'olivia.bennett@example.com', '+1 (415) 555-0198', '2847 Maple Street',
                  '94110', 'October 22, 2026', 'James Wilson', 'LAB-2026-08432'):
        assert value not in result['text']
    assert '118/76 mmHg' in result['text']
    assert 'Date of birth:' in result['text']
    assert 'Dr.' in result['text']
    assert 'Visit date:' in result['text']
    assert result['text'].count('[PATIENT_NAME_1]') == 2
    assert result['language']['code'] == 'en'
    for a, b in zip(result['entities'], result['entities'][1:]):
        assert a['end'] <= b['start']

@pytest.mark.parametrize('mode,token', [('redact','[REDACTED]'), ('blocks','█████')])
def test_other_modes(mode, token):
    assert anonymize('mail me at alice@example.com', mode)['text'] == f'mail me at {token}'

def test_unicode_offsets_and_html_are_plain_text():
    text = '🔒 <script>alert(1)</script> email: jane@example.com'
    result = anonymize(text)
    entity = next(e for e in result['entities'] if e['type'] == 'EMAIL')
    assert text[entity['start']:entity['end']] == 'jane@example.com'
    assert '<script>' in result['text']  # React must render as text, never as HTML.

def test_category_filter_and_empty_selection():
    assert anonymize('jane@example.com', enabled=[])['text'] == 'jane@example.com'
    assert anonymize('DOB: 2001/01/24')['counts']['DOB'] == 1

def test_local_pipeline_has_no_network(monkeypatch):
    def denied(*args, **kwargs):
        raise AssertionError('Runtime detection attempted outbound network access')
    monkeypatch.setattr(socket, 'connect', denied)
    assert anonymize(TEXT)['entities']

@pytest.mark.parametrize('body', [{}, {'text': ''}, {'text': 42}, {'text':'a'*100001},
    {'text': 'hello', 'enabled': ['INVALID']}, {'text':'hello', 'mode':'invalid'}])
def test_bad_requests(client, body):
    assert client.post('/api/anonymize', json=body).status_code == 400

def test_api_security_headers_and_csrf(client):
    r = client.post('/api/anonymize', json={'text':TEXT})
    assert r.status_code == 200
    assert 'Olivia Bennett' not in r.json['text']
    assert r.headers['Cache-Control'] == 'no-store'
    assert "connect-src 'self'" in r.headers['Content-Security-Policy']
    assert client.post('/api/anonymize', json={'text':TEXT}, headers={'Origin':'https://evil.example'}).status_code == 403
    assert client.post('/api/extract', headers={'Sec-Fetch-Site':'cross-site'}).status_code == 403

def test_utf8_txt_upload_and_unsupported_file(client):
    r = client.post('/api/extract', data={'file':(BytesIO(TEXT.encode()), 'notes.txt')})
    assert r.status_code == 200 and r.json['text'] == TEXT
    assert client.post('/api/extract', data={'file':(BytesIO(b'bad'), 'file.exe')}).status_code == 400
    assert client.post('/api/extract', data={'file':(BytesIO(b'\xff'), 'file.txt')}).status_code == 400
    assert client.post('/api/extract').status_code == 400

def test_docx_paragraph_and_table_extraction(client):
    doc = Document()
    doc.add_paragraph('Patient: Olivia Bennett')
    doc.add_table(rows=1, cols=1).cell(0,0).text = 'MRN: MR-12345'
    buffer = BytesIO(); doc.save(buffer); buffer.seek(0)
    r = client.post('/api/extract', data={'file':(buffer,'notes.docx')})
    assert r.status_code == 200
    assert 'Olivia Bennett' in r.json['text'] and 'MR-12345' in r.json['text']

def test_pdf_extraction_and_export_do_not_include_original(client):
    buffer = BytesIO(); pdf = canvas.Canvas(buffer)
    pdf.drawString(50,700, 'Patient: Olivia Bennett'); pdf.save(); buffer.seek(0)
    r = client.post('/api/extract', data={'file':(buffer,'notes.pdf')})
    assert r.status_code == 200 and 'Olivia Bennett' in r.json['text']
    clean = client.post('/api/anonymize', json={'text':TEXT}).json['text']
    result = client.post('/api/export/pdf', json={'text':clean})
    assert result.status_code == 200 and result.data.startswith(b'%PDF')
    text = '\n'.join(p.extract_text() for p in PdfReader(BytesIO(result.data)).pages)
    assert 'Olivia Bennett' not in text and 'PATIENT_NAME_1' in text

def test_corrupt_file_never_echoes_parser_data(client):
    result = client.post('/api/extract', data={'file':(BytesIO(b'%PDF secret-patient-name'), 'bad.pdf')})
    assert result.status_code == 422
    assert b'secret-patient-name' not in result.data

def test_library_value_error_is_not_a_public_validation_error(client, monkeypatch):
    import app.main as main
    def broken_parser(file):
        raise ValueError('secret-patient-name from malformed parser input')
    monkeypatch.setattr(main, 'extract_text', broken_parser)
    result = client.post('/api/extract', data={'file':(BytesIO(b'x'),'notes.pdf')})
    assert result.status_code == 422 and b'secret-patient-name' not in result.data

def test_language_is_deterministic_and_short_text_is_unknown():
    assert detect_language(TEXT) == detect_language(TEXT)
    assert detect_language('1234')['code'] == 'und'
    assert detect_language('Bonjour, ceci est un document qui contient des informations personnelles et des renseignements confidentiels.')['code'] == 'fr'

def test_image_ocr_path_and_missing_engine(client, monkeypatch):
    from PIL import Image
    import app.extraction as extraction
    monkeypatch.setattr(extraction, 'ocr_available', lambda: False)
    stream = BytesIO(); Image.new('RGB',(100,100),'white').save(stream,format='PNG'); stream.seek(0)
    result = client.post('/api/extract', data={'file':(stream,'image.png')})
    assert result.status_code == 400 and 'Tesseract' in result.json['error']

def test_upload_size_limit(client):
    r = client.post('/api/extract', data={'file':(BytesIO(b'a'*(10*1024*1024+1)), 'huge.txt')})
    assert r.status_code in (400,413)

def test_rebinding_host_is_rejected(client):
    assert client.get('/api/health', headers={'Host':'attacker.example'}).status_code == 400

@pytest.mark.parametrize('value', ['+91 98765 43210', '+44 20 7946 0958', '555-0198',
    '1234 5678 9012', 'ABCDE1234F', 'National ID: ZX-12345678', 'SSN: 482917364',
    'Lab accession number: LAB-123456', 'Sample ID S-123456', 'Barcode ID: BC987654'])
def test_international_and_contextual_identifiers(value):
    result = anonymize(value)
    assert result['entities']
    assert value not in result['text']

@pytest.mark.parametrize('extension', ['png','jpg','tiff'])
def test_image_extraction_uses_local_ocr(client, monkeypatch, extension):
    from PIL import Image
    import app.extraction as extraction
    monkeypatch.setattr(extraction, 'ocr_available', lambda: True)
    monkeypatch.setattr(extraction.pytesseract, 'image_to_string', lambda *a, **kw: 'Patient: Olivia Bennett')
    image = BytesIO(); Image.new('RGB',(100,100),'white').save(image, format='JPEG' if extension=='jpg' else extension.upper()); image.seek(0)
    result = client.post('/api/extract', data={'file':(image,'document.'+extension)})
    assert result.status_code == 200 and 'Olivia Bennett' in result.json['text']

def test_scanned_pdf_rasterizes_to_local_ocr(client, monkeypatch):
    import app.extraction as extraction
    monkeypatch.setattr(extraction, 'ocr_available', lambda: True)
    monkeypatch.setattr(extraction.pytesseract, 'image_to_string', lambda *a, **kw: 'MRN: MR-12345')
    stream = BytesIO(); pdf = canvas.Canvas(stream); pdf.showPage(); pdf.save(); stream.seek(0)
    result = client.post('/api/extract', data={'file':(stream,'scanned.pdf')})
    assert result.status_code == 200 and 'MR-12345' in result.json['text']
